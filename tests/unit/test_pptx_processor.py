import io

from pptx import Presentation

from backend.app.rag.ingestion.pptx_processor import extract_pptx_pages


def _make_pptx(slides: list[tuple[str, str | None]]) -> bytes:
    """slides: list of (bullet_text, notes_text | None)."""
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for bullet_text, notes_text in slides:
        slide = prs.slides.add_slide(layout)
        slide.placeholders[1].text_frame.text = bullet_text
        if notes_text is not None:
            slide.notes_slide.notes_text_frame.text = notes_text

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_one_page_per_slide_in_order():
    payload = _make_pptx([("First slide body", None), ("Second slide body", None)])

    pages = extract_pptx_pages(payload)

    assert [p for p, _ in pages] == [1, 2]
    assert "First slide body" in pages[0][1]
    assert "Second slide body" in pages[1][1]


def test_speaker_notes_are_included():
    payload = _make_pptx([("Bullet text", "This is the explanatory speaker note.")])

    _, text = extract_pptx_pages(payload)[0]

    assert "Bullet text" in text
    assert "Speaker notes: This is the explanatory speaker note." in text


def test_slide_without_notes_has_no_notes_section():
    payload = _make_pptx([("Just a bullet", None)])

    _, text = extract_pptx_pages(payload)[0]

    assert "Speaker notes" not in text
