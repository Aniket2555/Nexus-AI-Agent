import io

from pptx import Presentation


def extract_pptx_pages(payload: bytes) -> list[tuple[int, str]]:
    """One "page" per slide: shape text plus speaker notes.

    Notes are folded into the same page rather than dropped or split out — they're
    frequently where the actual explanatory content lives, versus the terse bullet
    text on the slide itself.
    """
    prs = Presentation(io.BytesIO(payload))
    pages: list[tuple[int, str]] = []
    for number, slide in enumerate(prs.slides, 1):
        parts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                parts.append(f"Speaker notes: {notes}")
        pages.append((number, "\n\n".join(parts)))
    return pages
